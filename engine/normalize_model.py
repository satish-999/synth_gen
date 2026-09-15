"""
normalize_model.py — validate and repair a SynthGen data-model workbook before
it is registered.

Every defect class below was observed in a real agent-produced model and caused
either a generation crash, a validation failure, or an infinite hang. All checks
are deterministic; no AI call is made.

CLI
    python normalize_model.py --in model.xlsx --out model_fixed.xlsx \
        [--report report.json] [--rows-hint 200000] [--dry-run]

Library
    from normalize_model import normalize
    result = normalize("model.xlsx", "model_fixed.xlsx", rows_hint=200000)
    result.ok            # True when nothing is BLOCKING
    result.changes       # list[Change]
    result.to_dict()     # JSON-ready, for the tool's review screen

Severity of each change:
    FIXED          repaired automatically, meaning preserved
    APPROXIMATED   repaired, but some intent was lost - show the user
    BLOCKING       cannot be repaired automatically - user must decide
"""
from __future__ import annotations

import argparse
import json
import re
from dataclasses import dataclass, field, asdict
from pathlib import Path

from openpyxl import load_workbook

# --------------------------------------------------------------------------- #
# what the engine actually supports
# --------------------------------------------------------------------------- #
GENERATORS = {
    "pattern", "faker", "choice", "numeric", "date", "date_offset", "sequence",
    "fk_lookup", "fk_lookup_jitter", "derived", "case",
}
RULE_TYPES = {"temporal_order", "bound", "conditional", "derived"}
NUMERIC_DISTS = {"uniform", "uniform_int", "lognormal"}
# functions pandas.eval can actually execute inside an expression
EVAL_SAFE = {"abs", "sqrt", "log", "exp", "sin", "cos", "tan"}

HEADERS = ["column_name", "data_type", "is_pk", "fk_ref", "fk_mode", "cardinality",
           "orphan_pct", "generator", "gen_params", "nullable_pct", "is_unique",
           "business_rule"]
# some workbooks use shorter header names for the same columns
HEADER_ALIASES = {
    "dtype": "data_type", "pk": "is_pk", "params": "gen_params",
    "unique": "is_unique", "nullable": "nullable_pct",
}

FIXED, APPROX, BLOCKING = "FIXED", "APPROXIMATED", "BLOCKING"


@dataclass
class Change:
    severity: str
    where: str          # "sheet.column" or "sheet"
    problem: str
    action: str
    before: str = ""
    after: str = ""


@dataclass
class Result:
    ok: bool = True
    changes: list = field(default_factory=list)
    tables: list = field(default_factory=list)
    views: list = field(default_factory=list)

    def add(self, severity, where, problem, action, before="", after=""):
        self.changes.append(Change(severity, where, problem, action,
                                   str(before)[:120], str(after)[:120]))
        if severity == BLOCKING:
            self.ok = False

    def counts(self):
        c = {FIXED: 0, APPROX: 0, BLOCKING: 0}
        for ch in self.changes:
            c[ch.severity] += 1
        return c

    def to_dict(self):
        return {"ok": self.ok, "counts": self.counts(), "tables": self.tables,
                "views": self.views, "changes": [asdict(c) for c in self.changes]}


# --------------------------------------------------------------------------- #
# helpers
# --------------------------------------------------------------------------- #
def _params(raw: str) -> dict:
    out = {}
    for part in str(raw or "").split(";"):
        if "=" in part:
            k, v = part.split("=", 1)
            out[k.strip()] = v.strip()
    return out


def _unparams(d: dict) -> str:
    return "; ".join(f"{k}={v}" for k, v in d.items())


def _pattern_capacity(pat: str) -> int:
    return (10 ** pat.count("#")) * (26 ** pat.count("?"))


def _header_map(ws) -> dict:
    """column name -> 1-based index, tolerating alias header names"""
    hm = {}
    for i, cell in enumerate(ws[1], start=1):
        name = str(cell.value or "").strip()
        hm[HEADER_ALIASES.get(name, name)] = i
    return hm


# --------------------------------------------------------------------------- #
# the normalizer
# --------------------------------------------------------------------------- #
def normalize(src: str | Path, dst: str | Path | None = None,
              rows_hint: int = 100_000) -> Result:
    """
    rows_hint: the largest row count a user might request. Unique columns whose
    generator cannot produce that many distinct values are widened, because that
    is what causes the engine to spin forever looking for a unique value.
    """
    res = Result()
    wb = load_workbook(src)

    # ---- object registry --------------------------------------------------- #
    if "_OBJECTS" not in wb.sheetnames:
        res.add(BLOCKING, "_OBJECTS", "registry sheet is missing",
                "add an _OBJECTS sheet listing every table and view")
        return res

    objs = {}
    for row in wb["_OBJECTS"].iter_rows(min_row=2, values_only=True):
        if row and row[0] and str(row[0]) != "workbook_version":
            objs[str(row[0])] = str(row[1] or "TABLE").upper()

    tables = [t for t, k in objs.items() if k == "TABLE"]
    views = [v for v, k in objs.items() if k == "VIEW"]
    res.tables, res.views = tables, views

    for name in list(objs):
        if name not in wb.sheetnames:
            res.add(BLOCKING, name, "listed in _OBJECTS but no sheet of that name",
                    "rename the sheet to match exactly, or remove the registry row")
            objs.pop(name)
            if name in tables: tables.remove(name)
            if name in views: views.remove(name)

    # ---- first pass: read every table ------------------------------------- #
    model = {}
    for t in tables:
        ws = wb[t]
        hm = _header_map(ws)
        missing = [h for h in ("column_name", "data_type", "generator") if h not in hm]
        if missing:
            res.add(BLOCKING, t, f"sheet is missing header column(s): {missing}",
                    "use the standard 12-column table sheet layout")
            continue
        cols = []
        for r in range(2, ws.max_row + 1):
            nm = ws.cell(row=r, column=hm["column_name"]).value
            if nm is None or str(nm).strip() == "":
                continue
            cols.append({"row": r, "name": str(nm).strip(),
                         **{h: (ws.cell(row=r, column=hm[h]).value if h in hm else None)
                            for h in HEADERS if h != "column_name"}})
        model[t] = {"ws": ws, "hm": hm, "cols": cols}

    pk_of, cols_of = {}, {}
    for t, m in model.items():
        pks = [c["name"] for c in m["cols"] if str(c.get("is_pk") or "").upper() == "Y"]
        pk_of[t] = pks
        cols_of[t] = {c["name"] for c in m["cols"]}
        if not pks:
            res.add(BLOCKING, t, "no primary key marked",
                    "set is_pk = Y on the key column (or columns)")

    def put(t, col, field_, value):
        m = model[t]
        if field_ in m["hm"]:
            m["ws"].cell(row=col["row"], column=m["hm"][field_]).value = value
            col[field_] = value
            return True
        return False

    # ---- second pass: repair ------------------------------------------------ #
    for t, m in model.items():
        for col in m["cols"]:
            where = f"{t}.{col['name']}"
            gen = str(col.get("generator") or "").strip()
            fk = str(col.get("fk_ref") or "").strip()
            params = _params(col.get("gen_params"))
            is_pk = str(col.get("is_pk") or "").upper() == "Y"

            # data_type labels such as "string (PK)"
            dt = str(col.get("data_type") or "").strip()
            # decimal(14,4), varchar(50), numeric(10,2) are valid type syntax and
            # must be preserved; only prose in brackets is stripped
            SQL_TYPE = re.compile(
                r"^(decimal|numeric|dec|varchar|char|nvarchar|nchar|float|"
                r"double|number)\s*\(\s*\d+\s*(,\s*\d+\s*)?\)$", re.I)
            if "(" in dt and not SQL_TYPE.match(dt):
                put(t, col, "data_type", dt.split("(")[0].strip())
                res.add(FIXED, where, "data_type carried an explanatory label",
                        "label stripped", dt, dt.split("(")[0].strip())

            # --- foreign keys ------------------------------------------------ #
            if fk:
                if "|" in fk or "," in fk:
                    res.add(BLOCKING, where,
                            "polymorphic foreign key: one column referencing more "
                            "than one parent table",
                            "the engine samples from exactly one parent. Merge the "
                            "parents into one table, or use a separate nullable FK "
                            "column per parent", fk)
                    continue
                if "." not in fk:
                    res.add(BLOCKING, where, f"fk_ref '{fk}' is not TABLE.COLUMN",
                            "write the parent table and its key column", fk)
                    continue
                pt, pc = fk.split(".", 1)
                if pt not in model:
                    res.add(BLOCKING, where, f"fk_ref points at unknown table '{pt}'",
                            "add the parent table or correct the reference", fk)
                elif pc not in cols_of.get(pt, set()):
                    res.add(BLOCKING, where, f"'{pc}' is not a column of {pt}",
                            "point at the parent's primary key", fk)
                elif pc not in pk_of.get(pt, []):
                    res.add(APPROX, where, f"fk_ref targets {pt}.{pc}, which is not a "
                            f"primary key", "left as written; verify it is unique", fk)

                # DEFECT: an FK column given its own generator produces keys that
                # reference nothing at all
                if gen:
                    put(t, col, "generator", None)
                    put(t, col, "gen_params", None)
                    res.add(FIXED, where,
                            "foreign-key column had its own value generator, so it "
                            "would have produced keys matching no parent row",
                            "generator removed; values now come from the parent",
                            f"{gen} ({_unparams(params)})", "(taken from parent)")
                    gen, params = "", {}

                mode = str(col.get("fk_mode") or "").strip().upper()
                if mode not in ("SIZING", "REFERENCE", ""):
                    put(t, col, "fk_mode", "REFERENCE")
                    res.add(FIXED, where, f"unknown fk_mode '{mode}'",
                            "set to REFERENCE", mode, "REFERENCE")
                    mode = "REFERENCE"
                if not mode:
                    put(t, col, "fk_mode", "REFERENCE")
                    res.add(FIXED, where, "fk_mode was blank",
                            "set to REFERENCE (parent-driven sizing not assumed)",
                            "", "REFERENCE")
                    mode = "REFERENCE"
                if mode == "SIZING":
                    card = str(col.get("cardinality") or "").strip()
                    if not re.match(r"^\s*\d+\s*,\s*\d+\s*,\s*(poisson\([\d.]+\)|\d+)\s*$",
                                    card):
                        put(t, col, "fk_mode", "REFERENCE")
                        put(t, col, "cardinality", None)
                        res.add(APPROX, where,
                                "SIZING relationship without a usable cardinality "
                                f"('{card}')",
                                "changed to REFERENCE; child volume now comes from "
                                "the row count instead of the parent", card, "REFERENCE")
                    else:
                        mn, mx = (int(x) for x in card.split(",")[:2])
                        if mn > mx:
                            put(t, col, "cardinality", f"{mx},{mn},{card.split(',')[2]}")
                            res.add(FIXED, where, "cardinality min greater than max",
                                    "bounds swapped", card)
                continue  # FK columns need no further generator checks

            # --- non-FK columns --------------------------------------------- #
            if not gen:
                res.add(BLOCKING, where, "no generator and no fk_ref",
                        "choose a generator, or make the column a foreign key")
                continue

            if gen not in GENERATORS:
                res.add(BLOCKING, where, f"generator '{gen}' does not exist",
                        f"supported: {', '.join(sorted(GENERATORS))}", gen)
                continue

            # nullable_pct given as a percentage instead of a fraction
            np_ = col.get("nullable_pct")
            if np_ not in (None, ""):
                try:
                    f = float(np_)
                    if f > 1:
                        put(t, col, "nullable_pct", round(f / 100, 4))
                        res.add(FIXED, where,
                                f"nullable_pct = {np_} would mean {f*100:.0f}% nulls",
                                "converted to a fraction", np_, round(f / 100, 4))
                        f = f / 100
                    if is_pk and f > 0:
                        put(t, col, "nullable_pct", 0)
                        res.add(FIXED, where,
                                "primary key allowed nulls, which fails validation",
                                "nullable_pct set to 0", np_, 0)
                except (TypeError, ValueError):
                    put(t, col, "nullable_pct", 0)
                    res.add(FIXED, where, f"nullable_pct '{np_}' is not a number",
                            "set to 0", np_, 0)

            dt_now = str(col.get("data_type") or "").strip().lower()
            def _numeric_choice():
                vs = [v.strip() for v in params.get("values", "").split(",") if v.strip()]
                if not vs:
                    return False
                try:
                    [float(v) for v in vs]
                    return True
                except ValueError:
                    return False

            TEXTY = {"pattern", "faker"} if gen != "choice" else (
                     set() if _numeric_choice() else {"choice"})
            if gen in TEXTY and dt_now in ("date", "timestamp", "int", "decimal"):
                put(t, col, "data_type", "string")
                res.add(FIXED, where,
                        f"column is typed '{dt_now}' but the {gen} generator "
                        f"produces text", "data_type set to string", dt_now, "string")
            if gen in ("date", "date_offset") and dt_now not in ("date", "timestamp", ""):
                put(t, col, "data_type", "date")
                res.add(FIXED, where,
                        f"column is typed '{dt_now}' but generates a date",
                        "data_type set to date", dt_now, "date")

            unique = str(col.get("is_unique") or "").upper() == "Y"
            if is_pk and not unique:
                put(t, col, "is_unique", "Y")
                res.add(FIXED, where, "primary key not marked unique",
                        "is_unique set to Y")
                unique = True

            # ---- per-generator checks -------------------------------------- #
            if gen == "pattern":
                pat = params.get("pattern", "")
                # prose leaking into the pattern corrupts every generated value
                m_ = re.match(r"^([#?A-Za-z0-9_\-/.]+)\s*[\(\[]", pat)
                if m_:
                    params["pattern"] = m_.group(1)
                    put(t, col, "gen_params", _unparams(params))
                    res.add(FIXED, where, "pattern contained explanatory text",
                            "text removed", pat, m_.group(1))
                    pat = m_.group(1)
                if not pat:
                    res.add(BLOCKING, where, "pattern generator without a pattern",
                            "add pattern=...")
                elif "#" not in pat and "?" not in pat:
                    res.add(BLOCKING, where,
                            f"pattern '{pat}' has no # or ? placeholders, so every "
                            f"row would get the same value", "add placeholders", pat)
                elif unique and _pattern_capacity(pat) < rows_hint:
                    before_cap = _pattern_capacity(pat)
                    runs = list(re.finditer(r"#+", pat))
                    if runs:
                        # widen the longest run of digits, wherever it sits in
                        # the pattern (handles 'IND-###-U#' and '###/#')
                        target = max(runs, key=lambda m: len(m.group()))
                        extra = 0
                        newpat = pat
                        while _pattern_capacity(newpat) < rows_hint and extra < 12:
                            extra += 1
                            newpat = (pat[: target.start()]
                                      + "#" * (len(target.group()) + extra)
                                      + pat[target.end():])
                        if newpat != pat:
                            params["pattern"] = newpat
                            put(t, col, "gen_params", _unparams(params))
                            res.add(FIXED, where,
                                    f"unique pattern '{pat}' can only produce "
                                    f"{before_cap:,} values; a larger run would hang "
                                    f"searching for a unique value",
                                    f"widened to {_pattern_capacity(newpat):,} values",
                                    pat, newpat)
                        else:
                            res.add(BLOCKING, where,
                                    f"unique pattern '{pat}' cannot be widened "
                                    f"automatically ({before_cap:,} values)",
                                    "add more # placeholders by hand", pat)
                    elif before_cap < 1000:
                        res.add(BLOCKING, where,
                                f"unique pattern '{pat}' can only produce "
                                f"{before_cap:,} values", "add # digits", pat)
                    else:
                        # letters only, e.g. 'CRP-???' = 17,576 values: fine for a
                        # dimension table, risky for a large one
                        res.add(APPROX, where,
                                f"unique pattern '{pat}' allows {before_cap:,} "
                                f"values; safe for a small table, but a run larger "
                                f"than that would hang",
                                "add a # digit if this table can grow", pat)

            elif gen == "choice":
                vals = [v for v in params.get("values", "").split(",") if v.strip()]
                if not vals:
                    res.add(BLOCKING, where, "choice generator without values",
                            "add values=a,b,c")
                else:
                    if any("(" in v for v in vals):
                        original = params["values"]
                        cleaned = ",".join(re.sub(r"\s*\([^)]*\)", "", v).strip()
                                           for v in vals)
                        params["values"] = cleaned
                        put(t, col, "gen_params", _unparams(params))
                        res.add(FIXED, where, "value list contained explanatory text",
                                "text removed", original, cleaned)
                        vals = cleaned.split(",")
                    if "weights" in params:
                        ws_ = [w for w in params["weights"].split(",") if w.strip()]
                        if len(ws_) != len(vals):
                            params.pop("weights")
                            put(t, col, "gen_params", _unparams(params))
                            res.add(FIXED, where,
                                    f"{len(vals)} values but {len(ws_)} weights, which "
                                    f"crashes the generator",
                                    "weights removed; values now equally likely")
                    if unique and len(vals) < rows_hint:
                        res.add(BLOCKING, where,
                                f"column is unique but choice offers only {len(vals)} "
                                f"values; the engine would hang", "use a pattern "
                                "generator, or drop the uniqueness requirement")

            elif gen == "faker":
                prov = params.get("provider", "")
                if not prov:
                    res.add(BLOCKING, where, "faker generator without a provider",
                            "add provider=name|company|email|address")
                elif prov in ("word", "words", "text", "sentence"):
                    res.add(APPROX, where,
                            f"faker provider '{prov}' produces random dictionary "
                            f"words, which look wrong in a business column",
                            "left as written - replace with a choice generator "
                            "listing the real domain values", prov)

            elif gen == "numeric":
                dist = params.get("dist", "uniform")
                if dist not in NUMERIC_DISTS:
                    params["dist"] = "uniform"
                    put(t, col, "gen_params", _unparams(params))
                    res.add(FIXED, where, f"unknown distribution '{dist}'",
                            "set to uniform", dist, "uniform")
                    dist = "uniform"
                if dist in ("uniform", "uniform_int"):
                    for k in ("min", "max"):
                        if k not in params:
                            res.add(BLOCKING, where,
                                    f"numeric {dist} without {k}", f"add {k}=")
                    try:
                        if float(params.get("min", 0)) > float(params.get("max", 0)):
                            params["min"], params["max"] = params["max"], params["min"]
                            put(t, col, "gen_params", _unparams(params))
                            res.add(FIXED, where, "min greater than max",
                                    "bounds swapped")
                    except ValueError:
                        res.add(BLOCKING, where, "min/max are not numbers",
                                "use numeric bounds")

            elif gen == "date":
                for k in ("start", "end"):
                    if k not in params:
                        res.add(BLOCKING, where, f"date generator without {k}",
                                f"add {k}=YYYY-MM-DD")

            elif gen == "date_offset":
                base = params.get("base")
                if not base:
                    res.add(BLOCKING, where, "date_offset without base",
                            "add base=<another date column>")
                elif base not in cols_of[t]:
                    res.add(BLOCKING, where,
                            f"date_offset base '{base}' is not a column of {t}",
                            "reference a date column on the same table", base)

            elif gen == "sequence":
                scope = params.get("scope")
                if scope and scope not in cols_of[t]:
                    res.add(BLOCKING, where,
                            f"sequence scope '{scope}' is not a column of {t}",
                            "scope must be the parent key column", scope)

            elif gen in ("fk_lookup", "fk_lookup_jitter"):
                srcp = params.get("source", "")
                if "." not in srcp:
                    res.add(BLOCKING, where, f"{gen} source '{srcp}' is not TABLE.COLUMN",
                            "write source=PARENT.column", srcp)
                else:
                    st, sc = srcp.split(".", 1)
                    if st not in model:
                        res.add(BLOCKING, where, f"source table '{st}' does not exist",
                                "correct the source", srcp)
                    elif sc not in cols_of.get(st, set()):
                        res.add(BLOCKING, where, f"'{sc}' is not a column of {st}",
                                "correct the source column", srcp)
                via = params.get("via")
                if via and ("+" in via or "-" in via):
                    res.add(BLOCKING, where,
                            f"composite via '{via}' asks for a multi-key or "
                            f"timestamp-matched join, which the engine cannot do",
                            "look the value up through a single key column", via)
                elif via and via not in cols_of[t]:
                    res.add(BLOCKING, where, f"via column '{via}' is not on {t}",
                            "name a key column present on this table", via)

            elif gen in ("derived", "case"):
                exprs = []
                if gen == "derived":
                    exprs = [("expr", params.get("expr", ""))]
                else:
                    exprs = [(k, v) for k, v in params.items()
                             if k.startswith("when")]
                for key, e in exprs:
                    if not e:
                        res.add(BLOCKING, where, f"{gen} generator without {key}",
                                f"add {key}=")
                        continue
                    if "^" in e:
                        params[key] = e.replace("^", "**")
                        put(t, col, "gen_params", _unparams(params))
                        res.add(FIXED, where,
                                "'^' means exclusive-or in the expression engine, "
                                "not exponentiation", "converted to '**'", e,
                                params[key])
                        e = params[key]
                    if "?" in e and ":" in e:
                        res.add(BLOCKING, where,
                                "ternary expression (a ? b : c) is not supported",
                                "use a case generator with when/then instead", e)
                        continue
                    if " contains " in e:
                        res.add(BLOCKING, where,
                                "'contains' is not an expression operator",
                                "compare with == against the exact value", e)
                        continue
                    for fn in re.findall(r"\b([a-z_][a-z0-9_]{2,})\s*\(", e):
                        if fn not in EVAL_SAFE:
                            res.add(BLOCKING, where,
                                    f"expression calls '{fn}()', which the engine "
                                    f"cannot execute",
                                    "replace with arithmetic, or generate the value "
                                    "directly with a numeric range", e)
                            break
                    for ident in re.findall(r"[A-Za-z_][A-Za-z0-9_]*", e):
                        if (ident not in cols_of[t] and ident not in EVAL_SAFE
                                and not ident.isupper()
                                and ident not in ("and", "or", "not", "True", "False")):
                            res.add(BLOCKING, where,
                                    f"expression references '{ident}', which is not a "
                                    f"column of {t}",
                                    "only columns on the same row can be used", e)
                            break

    # ---- views -------------------------------------------------------------- #
    for v in views:
        ws = wb[v]
        hdr, in_cols, vcols = {}, False, []
        for row in ws.iter_rows(values_only=True):
            if row and str(row[0] or "").strip() == "column_name":
                in_cols = True
                continue
            if not in_cols and row and row[0]:
                hdr[str(row[0]).strip()] = str(row[1] or "")
            elif in_cols and row and row[0]:
                vcols.append((str(row[0]), str(row[2] or "")))
        srcs = [s.strip() for s in hdr.get("source_objects", "").split(",") if s.strip()]
        if not srcs:
            res.add(BLOCKING, v, "view has no source_objects",
                    "name the table(s) the view is computed from")
        for s in srcs:
            if s not in model:
                res.add(BLOCKING, v, f"source object '{s}' is not a table in this model",
                        "correct the source", s)
        for name, der in vcols:
            if "rank(" in der or "over (" in der or "row_number(" in der:
                res.add(BLOCKING, f"{v}.{name}",
                        "window function in a view derivation",
                        "supported aggregates are count/sum/avg/min/max", der)
            elif der.startswith("round(") and re.search(
                    r"\b(count|sum|avg|min|max)\(", der[6:]):
                res.add(BLOCKING, f"{v}.{name}",
                        "aggregate nested inside round() in a post-aggregate column",
                        "reference the already-aggregated column instead", der)

    # ---- rules -------------------------------------------------------------- #
    if "_RULES" in wb.sheetnames:
        rws = wb["_RULES"]
        for r in range(2, rws.max_row + 1):
            rid = rws.cell(row=r, column=1).value
            if not rid:
                continue
            obj = str(rws.cell(row=r, column=2).value or "")
            rtype = str(rws.cell(row=r, column=3).value or "").strip()
            dfn = str(rws.cell(row=r, column=4).value or "")
            where = f"_RULES.{rid}"
            if rtype not in RULE_TYPES:
                continue                      # documentation row: engine ignores it
            if obj not in model:
                rws.cell(row=r, column=3).value = "note"
                res.add(FIXED, where, f"rule targets unknown object '{obj}'",
                        "retyped as a documentation note", rtype, "note")
                continue
            names = cols_of[obj]
            ok = True
            if rtype == "temporal_order":
                ok = bool(re.match(r"^\s*\w+(\s*<=\s*\w+)+\s*$", dfn))
            elif rtype == "bound":
                ok = bool(re.match(r"^\s*\w+\s*>=\s*\w+\s*$", dfn))
            elif rtype == "conditional":
                ok = bool(re.match(r"^when .+? then \w+ (NOT NULL|NULL)\s*$", dfn))
            elif rtype == "derived":
                ok = bool(re.match(r"^\s*\w+\s*=\s*[\w\s.*+\-/()]+$", dfn))
            refs = {x for x in re.findall(r"[A-Za-z_][A-Za-z0-9_]*", dfn)
                    if x not in ("when", "then", "NOT", "NULL", "in", "and", "or")}
            unknown = {x for x in refs if x not in names and not x.isupper()
                       and not x.replace(".", "").isdigit()}
            if not ok:
                rws.cell(row=r, column=3).value = "note"
                res.add(FIXED, where,
                        f"{rtype} rule is prose, not an expression the engine can "
                        f"evaluate", "retyped as a documentation note so it no "
                        "longer crashes the parser", dfn[:80], "note")
            elif unknown:
                rws.cell(row=r, column=3).value = "note"
                res.add(APPROX, where,
                        f"rule references {sorted(unknown)}, which are not columns "
                        f"of {obj}", "retyped as a documentation note; the rule is "
                        "no longer enforced", dfn[:80], "note")

    if dst:
        wb.save(dst)
    return res


# --------------------------------------------------------------------------- #
def main():
    ap = argparse.ArgumentParser(description="Normalize a SynthGen model workbook")
    ap.add_argument("--in", dest="src", required=True)
    ap.add_argument("--out", dest="dst")
    ap.add_argument("--report", dest="report")
    ap.add_argument("--rows-hint", type=int, default=100_000)
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()

    res = normalize(a.src, None if a.dry_run else (a.dst or a.src), a.rows_hint)
    c = res.counts()

    for ch in res.changes:
        print(f"  [{ch.severity:12}] {ch.where}")
        print(f"       {ch.problem}")
        print(f"       -> {ch.action}")
        if ch.before:
            print(f"       before: {ch.before}")
            if ch.after:
                print(f"       after:  {ch.after}")

    print(f"\n  {c[FIXED]} fixed, {c[APPROX]} approximated, {c[BLOCKING]} blocking")
    print("  RESULT:", "READY TO REGISTER" if res.ok
          else "BLOCKED - user decisions required")
    if a.report:
        Path(a.report).write_text(json.dumps(res.to_dict(), indent=2))
    raise SystemExit(0 if res.ok else 2)


if __name__ == "__main__":
    main()
