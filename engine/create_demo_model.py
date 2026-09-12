"""Bootstrap retail demo workbook for Phase 1 verification."""
from openpyxl import Workbook

HEADERS = [
    "column_name", "dtype", "pk", "fk_ref", "fk_mode", "cardinality",
    "orphan_pct", "generator", "params", "nullable_pct", "unique",
]


def write_table(ws, rows):
    ws.append(HEADERS)
    for row in rows:
        ws.append(row)


def main():
    wb = Workbook()
    wb.remove(wb.active)

    objects = wb.create_sheet("_OBJECTS")
    objects.append(["object_name", "object_type"])
    objects.append(["workbook_version", "1.1"])
    for name in ("category", "customer", "product", "sales_order", "order_item"):
        objects.append([name, "TABLE"])

    write_table(wb.create_sheet("category"), [
        ["category_id", "string", "Y", None, None, None, 0, "pattern", "pattern=CAT-##", 0, "Y"],
        ["category_name", "string", "N", None, None, None, 0, "choice",
         "values=Electronics,Clothing,Home,Grocery,Sports,Beauty", 0, "N"],
    ])

    write_table(wb.create_sheet("customer"), [
        ["customer_id", "string", "Y", None, None, None, 0, "pattern", "pattern=CUST-#####", 0, "Y"],
        ["customer_name", "string", "N", None, None, None, 0, "faker", "provider=name", 0, "N"],
        ["email", "string", "N", None, None, None, 0, "faker", "provider=email", 0, "N"],
    ])

    write_table(wb.create_sheet("product"), [
        ["product_id", "string", "Y", None, None, None, 0, "pattern", "pattern=PROD-#####", 0, "Y"],
        ["category_id", "string", "N", "category.category_id", "REFERENCE", None, 0, None, None, 0, "N"],
        ["product_name", "string", "N", None, None, None, 0, "faker", "provider=word", 0, "N"],
        ["unit_price", "float", "N", None, None, None, 0, "numeric", "dist=uniform; min=5; max=500; round=2", 0, "N"],
    ])

    write_table(wb.create_sheet("sales_order"), [
        ["order_id", "string", "Y", None, None, None, 0, "pattern", "pattern=ORD-######", 0, "Y"],
        ["customer_id", "string", "N", "customer.customer_id", "SIZING", "1,3,poisson(1.5)", 0.05, None, None, 0, "N"],
        ["order_date", "date", "N", None, None, None, 0, "date", "start=2024-01-01; end=2025-06-30", 0, "N"],
        ["status", "string", "N", None, None, None, 0, "choice",
         "values=PLACED,SHIPPED,DELIVERED,CANCELLED; weights=0.2,0.3,0.45,0.05", 0, "N"],
        ["ship_date", "date", "N", None, None, None, 0, "date_offset", "base=order_date; min_days=1; max_days=5", 0.1, "N"],
        ["delivery_date", "date", "N", None, None, None, 0, "date_offset", "base=ship_date; min_days=1; max_days=10", 0.2, "N"],
        ["delivery_status", "string", "N", None, None, None, 0, "case",
         "when1=status=='DELIVERED'; then1=DELIVERED; when2=status=='SHIPPED'; then2=IN_TRANSIT; else=PENDING", 0, "N"],
    ])

    write_table(wb.create_sheet("order_item"), [
        ["order_item_id", "string", "Y", None, None, None, 0, "pattern", "pattern=OI-#######", 0, "Y"],
        ["order_id", "string", "N", "sales_order.order_id", "SIZING", "1,5,poisson(2.5)", 0, None, None, 0, "N"],
        ["line_number", "int", "N", None, None, None, 0, "sequence", "scope=order_id; start=1", 0, "N"],
        ["product_id", "string", "N", "product.product_id", "REFERENCE", None, 0, None, None, 0, "N"],
        ["quantity", "int", "N", None, None, None, 0, "numeric", "dist=uniform_int; min=1; max=10", 0, "N"],
        ["unit_price", "float", "N", None, None, None, 0, "fk_lookup_jitter",
         "source=product.unit_price; jitter_pct=0.1; round=2", 0, "N"],
        ["line_total", "float", "N", None, None, None, 0, "derived", "expr=quantity * unit_price; round=2", 0, "N"],
    ])

    rules = wb.create_sheet("_RULES")
    rules.append(["rule_id", "object", "rule_type", "definition"])
    rules.append(["R1", "sales_order", "temporal_order", "order_date <= ship_date <= delivery_date"])
    rules.append(["R2", "sales_order", "conditional", "when status == 'PLACED' then delivery_date NULL"])
    rules.append(["R3", "sales_order", "bound", "delivery_date >= ship_date"])

    out = "data_model_demo.xlsx"
    wb.save(out)
    print(f"Created {out}")


if __name__ == "__main__":
    main()
