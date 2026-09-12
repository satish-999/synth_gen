"""
synthgen.py — workbook-driven relational synthetic data generator (framework v1.1).

Usage: python3 synthgen.py --model data_model.xlsx --config run_config.yaml
Stages: parse workbook -> resolve dependencies -> generate tables ->
        materialize views -> validate -> write CSV/XLSX + QA report + snapshot.
"""
import argparse, json, random, re, zlib
from datetime import datetime, timezone
from graphlib import TopologicalSorter
from pathlib import Path

import numpy as np
import pandas as pd
import yaml
from faker import Faker
from openpyxl import load_workbook

EXCEL_ROW_LIMIT = 1_048_576


# ---------------------------------------------------------------- parsing ---
def parse_params(raw):
    """'k1=v1; k2=a,b,c' -> dict with numbers coerced, lists split on comma."""
    out = {}
    if not raw:
        return out
    for part in str(raw).split(";"):
        if "=" not in part:
            continue
        k, v = part.split("=", 1)
        k, v = k.strip(), v.strip()
        if "," in v:
            out[k] = [_coerce(x.strip()) for x in v.split(",")]
        else:
            out[k] = _coerce(v)
    return out


def _coerce(v):
    for cast in (int, float):
        try:
            return cast(v)
        except ValueError:
            pass
    return v


def parse_workbook(path):
    wb = load_workbook(path, data_only=True)
    model = {"tables": {}, "views": {}, "rules": []}

    objects = {}
    for row in wb["_OBJECTS"].iter_rows(min_row=2, values_only=True):
        if row[0] and row[0] != "workbook_version":
            objects[row[0]] = row[1]

    for name, otype in objects.items():
        ws = wb[name]
        if otype == "TABLE":
            cols = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row[0]:
                    continue
                cols.append({
                    "name": row[0], "dtype": row[1], "pk": row[2] == "Y",
                    "fk_ref": row[3] or None, "fk_mode": row[4] or None,
                    "cardinality": row[5] or None,
                    "orphan_pct": float(row[6]) if row[6] not in (None, "") else 0.0,
                    "generator": row[7] or None, "params": parse_params(row[8]),
                    "nullable_pct": float(row[9]) if row[9] not in (None, "") else 0.0,
                    "unique": row[10] == "Y",
                })
            model["tables"][name] = cols
        else:
            spec, in_cols, cols = {}, False, []
            for row in ws.iter_rows(values_only=True):
                if row[0] == "column_name":
                    in_cols = True
                    continue
                if not in_cols and row[0]:
                    spec[row[0]] = row[1]
                elif in_cols and row[0]:
                    cols.append({"name": row[0], "dtype": row[1], "derivation": row[2]})
            spec["columns"] = cols
            model["views"][name] = spec

    for row in wb["_RULES"].iter_rows(min_row=2, values_only=True):
        if row[0]:
            model["rules"].append({"rule_id": row[0], "object": row[1],
                                   "rule_type": row[2], "definition": row[3]})
    return model


def validate_model(model):
    errors = []
    all_pk = {f"{t}.{c['name']}" for t, cols in model["tables"].items()
              for c in cols if c["pk"]}
    for t, cols in model["tables"].items():
        if not any(c["pk"] for c in cols):
            errors.append(f"{t}: no primary key defined")
        for c in cols:
            if c["fk_ref"] and c["fk_ref"] not in all_pk:
                errors.append(f"{t}.{c['name']}: fk_ref '{c['fk_ref']}' does not "
                              f"resolve to a PK column")
            if c["fk_mode"] == "SIZING" and not c["cardinality"]:
                errors.append(f"{t}.{c['name']}: SIZING fk needs cardinality")
            if not c["fk_ref"] and not c["generator"]:
                errors.append(f"{t}.{c['name']}: needs a generator or fk_ref")
    for v, spec in model["views"].items():
        for src in [s.strip() for s in spec["source_objects"].split(",")]:
            if src not in model["tables"]:
                errors.append(f"{v}: source object '{src}' not in workbook")
    if errors:
        raise SystemExit("MODEL VALIDATION FAILED:\n  - " + "\n  - ".join(errors))


# ------------------------------------------------------------- generation ---
class Engine:
    def __init__(self, model, cfg):
        self.m, self.cfg = model, cfg
        seed = cfg.get("seed", 0)
        random.seed(seed); np.random.seed(seed)
        self.fake = Faker(cfg.get("locale", "en_US"))
        self.fake.seed_instance(seed)
        self.frames = {}

    # dependency order: FK parents before children, views last
    def order(self):
        deps = {t: set() for t in self.m["tables"]}
        for t, cols in self.m["tables"].items():
            for c in cols:
                if c["fk_ref"]:
                    parent = c["fk_ref"].split(".")[0]
                    if parent != t:
                        deps[t].add(parent)
        tables = list(TopologicalSorter(deps).static_order())
        return tables, list(self.m["views"].keys())

    # ---- value generators ------------------------------------------------
    def gen_pattern(self, p, n):
        def one():
            return "".join(random.choice("0123456789") if ch == "#"
                           else random.choice("ABCDEFGHIJKLMNOPQRSTUVWXYZ") if ch == "?"
                           else ch for ch in str(p["pattern"]))
        return [one() for _ in range(n)]

    def gen_choice(self, p, n):
        vals = p["values"] if isinstance(p["values"], list) else [p["values"]]
        w = p.get("weights")
        w = w if isinstance(w, list) or w is None else [w]
        return random.choices(vals, weights=w, k=n)

    def gen_faker(self, p, n):
        return [getattr(self.fake, p["provider"])() for _ in range(n)]

    def gen_numeric(self, p, n):
        d = p.get("dist", "uniform")
        if d == "uniform_int":
            v = np.random.randint(p["min"], p["max"] + 1, n).astype(float)
        elif d == "lognormal":
            v = np.clip(np.random.lognormal(p["mean"], p["sigma"], n),
                        p.get("min", 0), p.get("max", np.inf))
        else:
            v = np.random.uniform(p["min"], p["max"], n)
        r = int(p.get("round", 0))
        v = np.round(v, r)
        return v.astype(int) if r == 0 and d == "uniform_int" else v

    def gen_date(self, p, n):
        s, e = pd.Timestamp(str(p["start"])), pd.Timestamp(str(p["end"]))
        return s + pd.to_timedelta(np.random.randint(0, (e - s).days + 1, n), unit="D")

    # ---- sizing allocation -------------------------------------------------
    def child_counts(self, card, orphan_pct, n_parents):
        mn, mx, dist = card if isinstance(card, list) else str(card).split(",")
        mn, mx = int(mn), int(mx)
        m = re.match(r"poisson\((.+)\)", str(dist).strip())
        counts = (np.random.poisson(float(m.group(1)), n_parents) if m
                  else np.random.randint(mn, mx + 1, n_parents))
        counts = np.clip(counts, mn, mx)
        counts[np.random.rand(n_parents) < orphan_pct] = 0
        return counts

    def _reseed(self, name):
        # per-table RNG stream: values depend on (seed, table) only — never on
        # which other tables are in the run. zlib.crc32 is stable across
        # processes (unlike hash() for str).
        base = int(self.cfg.get("seed", 0))
        s = (base * 1_000_003 + zlib.crc32(name.encode())) % (2**32)
        random.seed(s); np.random.seed(s); self.fake.seed_instance(s)

    def generate_table(self, name):
        self._reseed(name)
        cols = self.m["tables"][name]
        sizing = next((c for c in cols if c["fk_mode"] == "SIZING"), None)

        target = self.cfg["targets"][name]
        if sizing and target.get("rows") == "derived":
            pt, pc = sizing["fk_ref"].split(".")
            parents = self.frames[pt][pc].values
            counts = self.child_counts(sizing["cardinality"], sizing["orphan_pct"],
                                       len(parents))
            n = int(counts.sum())
            df = pd.DataFrame({sizing["name"]: np.repeat(parents, counts)})
        else:
            n = int(target["rows"])
            df = pd.DataFrame(index=range(n))

        deferred_case = []
        for c in cols:
            cn = c["name"]
            if cn in df.columns:
                continue
            p = c["params"]
            if c["fk_ref"]:                                     # REFERENCE fk
                pt, pc = c["fk_ref"].split(".")
                df[cn] = np.random.choice(self.frames[pt][pc].values, n)
            elif c["generator"] == "sequence":
                df[cn] = df.groupby(p["scope"]).cumcount() + int(p.get("start", 1))
            elif c["generator"] == "date_offset":
                base = pd.to_datetime(df[p["base"]])
                off = np.random.randint(int(p["min_days"]), int(p["max_days"]) + 1, n)
                df[cn] = base + pd.to_timedelta(off, unit="D")
            elif c["generator"] == "fk_lookup":
                st, sc = p["source"].split(".")
                fk = p.get("via") or next(x["name"] for x in cols
                          if (x["fk_ref"] or "").startswith(st + "."))
                pk = next(x["name"] for x in self.m["tables"][st] if x["pk"])
                df[cn] = df[fk].map(self.frames[st].set_index(pk)[sc])
            elif c["generator"] == "fk_lookup_jitter":
                st, sc = p["source"].split(".")
                fk = next(x["name"] for x in cols
                          if (x["fk_ref"] or "").startswith(st + "."))
                pk = next(x["name"] for x in self.m["tables"][st] if x["pk"])
                lut = self.frames[st].set_index(pk)[sc]
                jit = 1 + np.random.uniform(-p["jitter_pct"], p["jitter_pct"], n)
                mult = float(p.get("multiplier", 1))      # optional premium/discount
                df[cn] = np.round(df[fk].map(lut).astype(float) * jit * mult,
                                  int(p.get("round", 2)))
            elif c["generator"] == "case":
                deferred_case.append(c)          # evaluate AFTER rules
                df[cn] = ""                        # placeholder keeps column order
            elif c["generator"] == "derived":
                df[cn] = np.round(pd.Series(df.eval(p["expr"])).astype(float), int(p.get("round", 2)))
            else:
                fn = getattr(self, f"gen_{c['generator']}")
                vals = list(fn(p, n))
                if c["unique"]:
                    seen, out = set(), []
                    for v in vals:
                        tries = 0
                        while v in seen:
                            v = fn(p, 1)[0]
                            tries += 1
                            if tries > 10000:
                                raise ValueError(
                                    f"Cannot generate {n:,} unique values for "
                                    f"{name}.{cn}: the '{c['generator']}' generator "
                                    f"({p}) cannot produce enough distinct values. "
                                    f"Widen the pattern (add more # digits) or "
                                    f"request fewer rows.")
                        seen.add(v); out.append(v)
                    vals = out
                df[cn] = vals
            if c["nullable_pct"]:
                df.loc[np.random.rand(n) < c["nullable_pct"], cn] = (pd.NaT if "date" in str(c["dtype"]) else None)

        self.apply_rules(name, df)
        for c in deferred_case:                   # case logic sees final, rule-adjusted values
            p, conds, vals, i = c["params"], [], [], 1
            while f"when{i}" in p:
                conds.append(df.eval(p[f"when{i}"], engine="python"))
                vals.append(str(p[f"then{i}"])); i += 1
            df[c["name"]] = np.select(conds, vals, default=str(p.get("else", "")))
        self.frames[name] = df[[c["name"] for c in cols]]   # model-declared order

    # ---- rules -------------------------------------------------------------
    def apply_rules(self, name, df):
        for r in [x for x in self.m["rules"] if x["object"] == name]:
            d = r["definition"]
            if r["rule_type"] == "temporal_order":
                cols = [c.strip() for c in d.split("<=")]
                present = [c for c in cols if c in df.columns]
                vals = df[present].values.astype("datetime64[ns]")
                mask = ~pd.isna(vals).any(axis=1)          # keep nulls untouched
                vals[mask] = np.sort(vals[mask], axis=1)
                for i, c in enumerate(present):
                    col = pd.Series(vals[:, i], index=df.index)
                    df[c] = col.where(~df[c].isna(), pd.NaT)
            elif r["rule_type"] == "conditional":
                m = re.match(r"when (.+?) then (\w+) (NOT NULL|NULL)", d)
                cond, target, action = m.groups()
                mask = df.eval(cond)
                if action == "NULL":
                    df.loc[mask, target] = pd.NaT
                else:
                    pool = df.loc[~mask, target].dropna()
                    need = mask & df[target].isna()
                    if need.any() and len(pool):
                        df.loc[need, target] = pool.sample(
                            int(need.sum()), replace=True).values
            elif r["rule_type"] == "bound":
                m = re.match(r"(\w+)\s*>=\s*(\w+)", d)
                col, floor = m.groups()
                viol = df[col].notna() & (df[col] < df[floor])
                if viol.any():
                    bump = pd.to_timedelta(
                        np.random.randint(1, 6, int(viol.sum())), unit="D")
                    df.loc[viol, col] = df.loc[viol, floor] + bump
            elif r["rule_type"] == "derived":
                m = re.match(r"(\w+)\s*=\s*(.+)", d)
                df[m.group(1)] = np.round(df.eval(m.group(2)), 2)

    # ---- views: computed, never fabricated -----------------------------------
    def materialize_view(self, name):
        spec = self.m["views"][name]
        sources = [s.strip() for s in spec["source_objects"].split(",")]
        joined = {sources[0]}
        df = self.frames[sources[0]].copy()
        for j in [x.strip() for x in (spec.get("join_logic") or "").split(";") if x.strip()]:
            left, right = [s.strip() for s in j.split("=")]
            lt, lc = left.split("."); rt, rc = right.split(".")
            new_t, new_c, cur_c = (rt, rc, lc) if lt in joined else (lt, lc, rc)
            df = df.merge(self.frames[new_t], left_on=cur_c, right_on=new_c,
                          how="inner", suffixes=("", "_r"))
            joined.add(new_t)
        if spec.get("filter_logic"):
            df = df.query(str(spec["filter_logic"]))

        gb = [c.strip() for c in (spec.get("group_by") or "").split(",") if c.strip()]
        AGG = re.compile(r"^(count|sum|avg|min|max)\(")
        agg_cols = [c for c in spec["columns"] if AGG.match(c["derivation"])]
        post_cols = [c for c in spec["columns"]
                     if c["derivation"].startswith("round(")]

        if gb:
            rows = []
            for key, g in df.groupby(gb):
                key = key if isinstance(key, tuple) else (key,)
                rec = dict(zip(gb, key))
                for c in agg_cols:
                    rec[c["name"]] = eval_agg(c["derivation"], g)
                rows.append(rec)
            out = pd.DataFrame(rows)
        else:
            out = df
        want = [c["name"] for c in spec["columns"]]
        if len(out) == 0:                      # filter removed every row
            self.frames[name] = pd.DataFrame(columns=want)
            return
        for c in post_cols:
            inner = re.match(r"round\((.+),\s*(\d+)\)", c["derivation"])
            out[c["name"]] = np.round(out.eval(inner.group(1)),
                                      int(inner.group(2)))
        self.frames[name] = out[[c for c in want if c in out.columns]]


def eval_agg(expr, g):
    m = re.match(r"count\((\w+)\)", expr)
    if m:
        return int(g[m.group(1)].count())
    m = re.match(r"sum\(case when (.+?) then 1 else 0 end\)", expr)
    if m:
        return int(g.eval(m.group(1)).sum())
    m = re.match(r"(sum|avg|min|max)\((\w+)\)", expr)
    if m:
        fn, col = m.groups()
        return getattr(g[col], {"avg": "mean"}.get(fn, fn))()
    raise ValueError(f"unsupported aggregate: {expr}")


# -------------------------------------------------------------- validation ---
def validate_output(model, frames):
    report, ok = {}, True
    for t, cols in model["tables"].items():
        if t not in frames:
            continue
        checks, df = [], frames[t]
        pks = [c["name"] for c in cols if c["pk"]]
        dup = int(df.duplicated(pks).sum())
        checks.append({"check": "pk_unique", "pass": dup == 0, "duplicates": dup})
        for c in cols:
            if c["fk_ref"]:
                pt, pc = c["fk_ref"].split(".")
                orphans = int((~df[c["name"]].dropna()
                               .isin(frames[pt][pc])).sum())
                checks.append({"check": f"fk_{c['name']}", "pass": orphans == 0,
                               "orphans": orphans})
        report[t] = checks
        ok &= all(c["pass"] for c in checks)

    # view consistency: totals reconcile against base data
    for v in model["views"]:
        if v in frames and "total_lines" in frames[v].columns:
            base = frames["PO_LINE"].merge(
                frames["PURCHASE_ORDER"][["po_number", "status"]], on="po_number")
            expected = int((base["status"] != "CANCELLED").sum())
            actual = int(frames[v]["total_lines"].sum())
            match = expected == actual
            report[v] = [{"check": "view_reconciles_base", "pass": match,
                          "expected": expected, "actual": actual}]
            ok &= match
    return ok, report


# --------------------------------------------------------------------- main ---
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--model", required=True)
    ap.add_argument("--config", required=True)
    args = ap.parse_args()

    cfg = yaml.safe_load(Path(args.config).read_text())
    model = parse_workbook(args.model)
    validate_model(model)

    run_id = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    out = Path(cfg["output"]["path"]); out.mkdir(parents=True, exist_ok=True)
    (out / f"model_snapshot_{run_id}.json").write_text(json.dumps(model, indent=2, default=str))

    eng = Engine(model, cfg)

    # load previously generated data as fixed reference inputs (not regenerated)
    for name, ref in cfg.get("references", {}).items():
        df = pd.read_csv(ref["path"])
        for c in model["tables"].get(name, []):
            if c["dtype"] == "date" and c["name"] in df.columns:
                df[c["name"]] = pd.to_datetime(df[c["name"]])
        eng.frames[name] = df
        print(f"  ref   {name}: {len(df):,} rows loaded from {ref['path']}")

    # every FK parent of a target must be a target or a reference
    for t in cfg["targets"]:
        for c in model["tables"].get(t, []):
            if c["fk_ref"]:
                parent = c["fk_ref"].split(".")[0]
                if parent not in cfg["targets"] and parent not in eng.frames:
                    raise SystemExit(
                        f"CONFIG ERROR: {t} needs parent '{parent}' — add it to "
                        f"targets or provide it under references.")

    tables, views = eng.order()
    for t in tables:
        if t in cfg["targets"] and t not in eng.frames:
            eng.generate_table(t)
            print(f"  table {t}: {len(eng.frames[t]):,} rows")
    for v in views:
        srcs = [s.strip() for s in model["views"][v]["source_objects"].split(",")]
        if all(s in eng.frames for s in srcs) and (
                v in cfg["targets"] or not cfg.get("references")):
            eng.materialize_view(v)
            print(f"  view  {v}: {len(eng.frames[v]):,} rows (computed)")

    ok, report = validate_output(model, eng.frames)
    (out / "qa_report.json").write_text(json.dumps(
        {"run_id": run_id, "seed": cfg.get("seed"), "passed": ok,
         "objects": report}, indent=2))
    if not ok:
        raise SystemExit("VALIDATION FAILED — no data written. See qa_report.json")

    written = {n: d for n, d in eng.frames.items()
               if n not in cfg.get("references", {})}
    fmts = cfg["output"]["format"]
    for name, df in written.items():
        w = df.copy()
        for c in w.columns:                       # dates as ISO strings for CSV
            if pd.api.types.is_datetime64_any_dtype(w[c]):
                w[c] = w[c].dt.strftime("%Y-%m-%d")
        if "csv" in fmts:
            w.to_csv(out / f"{name}.csv", index=False)
    if "xlsx" in fmts:
        big = [n for n, d in written.items() if len(d) > EXCEL_ROW_LIMIT]
        if big:
            print(f"  ! skipped xlsx for oversized objects: {big}")
        with pd.ExcelWriter(out / "synthetic_data.xlsx") as xw:
            for name, df in written.items():
                if len(df) <= EXCEL_ROW_LIMIT:
                    w = df.copy()
                    for c in w.columns:
                        if pd.api.types.is_datetime64_any_dtype(w[c]):
                            w[c] = w[c].dt.strftime("%Y-%m-%d")
                    w.to_excel(xw, sheet_name=name[:31], index=False)
    print(f"  QA: {'PASS' if ok else 'FAIL'} -> {out}")


if __name__ == "__main__":
    main()
